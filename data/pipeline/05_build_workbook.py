"""
05_build_workbook.py -- screening workbook (sheets Criteria, Rules, Screening, PRISMA_Summary, Check).

The decisions come from pipeline_output/03_final_decisions.csv, written by the screening rules
(pipeline/02_screen_records.py, rules in pipeline/screening_rules.yaml). For each record the workbook gives the
decision, the exclusion reason, the rule that decided it and the matched terms. The Check sheet holds the blinded
reading of a random sample (pipeline/03_check_screening.py) once it has been scored. Publisher abstracts and
keywords are not written to the workbook, because the database exports cannot be redistributed.

Usage: python -I pipeline/05_build_workbook.py   (run from data/, after step 04)
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

IN_PATH = "pipeline_output/03_final_decisions.csv"
OUT_PATH = "pipeline_output/R01_PRISMA_screening_FINAL.xlsx"
CHECK_RESULTS = "check/screening_check_results.txt"

df = pd.read_csv(IN_PATH, encoding="utf-8-sig")
n_rows = len(df)
version, sha = df["rules_version"].iloc[0], df["rules_sha256"].iloc[0]
recon = pd.read_csv("pipeline_output/04_citation_reconciliation.csv", encoding="utf-8-sig")
n_cited, n_found = len(recon), int(recon["found_in_search"].sum())
HEAD = PatternFill("solid", fgColor="4472C4")
wb = Workbook()


def table(ws, rows, widths, header=True):
    for i, vals in enumerate(rows):
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i + 1, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if header and i == 0:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ---------------- Criteria ----------------
ws = wb.active
ws.title = "Criteria"
table(ws, [
    ("Criterion", "Meaning (protocol/eligibility_criteria.md)"),
    ("I3", "A stochastic model is developed or applied"),
    ("I4", "Genetics, evolution, plant or animal breeding, or gene-regulatory-network dynamics is the main subject"),
    ("I5", "Title and abstract are sufficient for a decision"),
    ("E1", "The stochastic model is applied outside these fields"),
    ("E2", "Genetics or evolution is mentioned only in passing"),
    ("E3", "The method is deterministic"),
    ("E5", "No usable abstract"),
    ("E6", "Conference abstract, erratum, retraction notice or another non-substantive document"),
    ("", ""),
    ("Exclusion reason", "Meaning"),
    ("not_genetics_domain", "Stochastic model applied outside genetics, evolution, breeding and gene regulation (E1)"),
    ("keyword_coincidence", "Search terms matched by another sense of the words, such as a genetic algorithm (E1, E2)"),
    ("not_stochastic_framework", "No stochastic model (E3)"),
    ("wrong_document_type", "Not a research or review article (E6)"),
    ("insufficient_information", "Abstract missing or too short to decide (E5, I5); kept as Maybe"),
], (26, 100))

# ---------------- Rules ----------------
ws = wb.create_sheet("Rules")
table(ws, [
    ("Rule", "Applies when", "Decision"),
    ("E6", "document type is excluded", "Exclude (wrong_document_type)"),
    ("E5", "no abstract", "Maybe (insufficient_information)"),
    ("E3", "no stochastic term", "Exclude (not_stochastic_framework)"),
    ("E3?", "no stochastic term but a strong population-genetic subject", "Maybe"),
    ("E1a", "no genetics, evolution, breeding or gene-regulation term", "Exclude (not_genetics_domain)"),
    ("E1c", "an off-domain subject term in the title and no strong domain term there",
     "Exclude (not_genetics_domain or keyword_coincidence)"),
    ("E1b", "off-domain score at least 2 and not below the biological domain score",
     "Exclude (not_genetics_domain or keyword_coincidence)"),
    ("E1b?", "as E1b, but the biological domain score is 4 or more", "Maybe"),
    ("I", "a stochastic model and enough domain evidence", "Include"),
    ("E2s / E2d", "the stochastic model is not named, or the domain evidence is below the threshold", "Maybe"),
    ("", "", ""),
    ("Rules file", f"pipeline/screening_rules.yaml, version {version}", f"SHA-256 {sha}"),
], (12, 80, 45))

# ---------------- Screening ----------------
ws = wb.create_sheet("Screening")
cols = [("ID", None, 6), ("Title", "title", 50), ("Authors", "authors", 22), ("Year", "year", 6),
        ("Source", "source_title", 22), ("DOI", "doi", 22), ("Database", "database", 13),
        ("Document_Type", "doc_type", 12), ("Decision", "final_decision", 11),
        ("Exclusion_Reason", "final_exclusion_reason", 22), ("Rule", "screen_rule", 7), ("Notes", "final_notes", 50),
        ("Domain_score", "domain_score", 8), ("Biological_domain_score", "domain_bio_score", 8),
        ("Off_domain_score", "off_domain_score", 8), ("Stochastic_terms", "stochastic_terms", 30),
        ("Domain_terms", "domain_terms", 30), ("Off_domain_terms", "off_domain_terms", 25)]
fill = {"Include": "E2EFDA", "Exclude": "FCE4E4", "Maybe": "FFF2CC"}
rows = [[c[0] for c in cols]]
for i, r in df.iterrows():
    rows.append([i + 1] + ["" if pd.isna(r[k]) else r[k] for _, k, _ in cols[1:]])
table(ws, rows, [c[2] for c in cols])
for r in range(2, n_rows + 2):
    c = ws.cell(row=r, column=9)
    if c.value in fill:
        c.fill = PatternFill("solid", fgColor=fill[c.value])
    c.font = Font(bold=True)
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n_rows + 1}"

# ---------------- PRISMA_Summary ----------------
ws = wb.create_sheet("PRISMA_Summary")
counts = df["final_decision"].value_counts()
reasons = df.loc[df["final_decision"] == "Exclude", "final_exclusion_reason"].value_counts()
summary = [("Stage", "Count", "Source"),
           ("Records identified - Scopus", 974, "data/raw/scopus/ (export of 26 July 2026)"),
           ("Records identified - Web of Science", 721, "data/raw/WoS/ (export of 26 July 2026)"),
           ("Duplicate records removed", 1695 - n_rows, "pipeline/01_dedupe.py"),
           ("Records screened (title and abstract)", n_rows, f"pipeline/02_screen_records.py, rules {version}"),
           ("  Include", int(counts.get("Include", 0)), ""),
           ("  Maybe (title and abstract not sufficient)", int(counts.get("Maybe", 0)), ""),
           ("  Exclude", int(counts.get("Exclude", 0)), "")]
summary += [(f"    {k}", int(v), "") for k, v in reasons.items()]
summary += [("Works cited in the manuscript", n_cited, "cited_references.csv"),
            ("  also retrieved by the search", n_found, "pipeline/04_reconcile_citations.py")]
table(ws, summary, (50, 10, 60))

# ---------------- Check ----------------
ws = wb.create_sheet("Check")
lines = open(CHECK_RESULTS, encoding="utf-8").read().splitlines() if os.path.exists(CHECK_RESULTS) else [
    "The blinded check sample has not been scored yet (pipeline/03_check_screening.py sample, then score)."]
table(ws, [(l,) for l in lines], (110,), header=False)

wb.save(OUT_PATH)
print(f"Saved {OUT_PATH}: {n_rows} records, rules {version}")
print(counts.to_string())
